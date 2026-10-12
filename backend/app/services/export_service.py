"""Excel 导出：openpyxl 只写模式，先落临时文件再流式回传，避免全量数据驻留内存。"""

import os
import tempfile
from datetime import datetime
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import Select
from sqlalchemy.orm import Session, selectinload

from app.models.enums import ORDER_STATUS_LABELS, ProductStatus
from app.models.order import Order
from app.models.product import Product

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="0052D9")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

ORDER_HEADERS: list[tuple[str, int]] = [
    ("订单号", 22),
    ("客户姓名", 14),
    ("联系电话", 16),
    ("收货地址", 42),
    ("商品名称", 26),
    ("规格", 16),
    ("单价", 12),
    ("数量", 10),
    ("小计", 12),
    ("订单金额", 12),
    ("订单状态", 12),
    ("下单时间", 22),
    ("来源链接备注", 22),
    ("客户备注", 22),
]

PRODUCT_HEADERS: list[tuple[str, int]] = [
    ("商品ID", 10),
    ("商品名称", 28),
    ("SKU", 18),
    ("所属分类", 16),
    ("单价", 12),
    ("单位", 8),
    ("规格", 40),
    ("状态", 10),
    ("商品描述", 46),
    ("创建时间", 22),
]


def _new_sheet(wb: Workbook, title: str, headers: list[tuple[str, int]]):
    ws = wb.create_sheet(title)
    cells: list[WriteOnlyCell] = []
    for index, (name, width) in enumerate(headers, start=1):
        cell = WriteOnlyCell(ws, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cells.append(cell)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.append(cells)
    ws.freeze_panes = "A2"
    return ws


def _make_temp_path(prefix: str) -> str:
    handle, path = tempfile.mkstemp(prefix=prefix, suffix=".xlsx")
    os.close(handle)
    return path


def _fmt_time(value: Optional[datetime]) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if value else ""


def _fmt_variants(product: Product) -> str:
    """规格压成一行，便于在 Excel 里筛选与对照；下架规格显式标记。"""
    parts = []
    for item in product.variants:
        suffix = "" if item.status == ProductStatus.ON_SALE.value else "（已下架）"
        parts.append(f"{item.name} ¥{float(item.price):.2f}{suffix}")
    return "；".join(parts)


def export_orders(db: Session, stmt: Select) -> str:
    """订单按明细行展开导出，一行一个商品。"""
    path = _make_temp_path("orders_")
    wb = Workbook(write_only=True)
    ws = _new_sheet(wb, "订单数据", ORDER_HEADERS)

    # selectinload 以独立查询批量取明细，不会产生行膨胀，因此无需 unique()；
    # 而 unique() 与 yield_per() 互斥，去掉后即可边读边写，内存占用恒定。
    query = db.execute(stmt.options(selectinload(Order.items))).scalars()
    for order in query.yield_per(200):
        status_label = ORDER_STATUS_LABELS.get(order.status, order.status)
        base: tuple[Any, ...] = (
            order.order_no,
            order.customer_name,
            order.customer_phone,
            order.customer_address,
        )
        tail: tuple[Any, ...] = (
            float(order.total_amount),
            status_label,
            _fmt_time(order.created_at),
            order.link_remark or "",
            order.remark or "",
        )
        items = order.items or []
        if not items:
            ws.append(list(base) + ["", "", "", "", "", *tail])
            continue
        for item in items:
            ws.append(
                list(base)
                + [
                    item.product_name,
                    item.variant_name or "",
                    float(item.price),
                    item.quantity,
                    float(item.subtotal),
                    *tail,
                ]
            )

    wb.save(path)
    return path


def export_products(db: Session, products: Iterable[Product]) -> str:
    path = _make_temp_path("products_")
    wb = Workbook(write_only=True)
    ws = _new_sheet(wb, "商品数据", PRODUCT_HEADERS)

    status_labels = {
        ProductStatus.ON_SALE.value: "上架",
        ProductStatus.OFF_SHELF.value: "下架",
    }
    for product in products:
        ws.append(
            [
                product.id,
                product.name,
                product.sku or "",
                product.category.name if product.category else "",
                float(product.price),
                product.unit,
                _fmt_variants(product),
                status_labels.get(product.status, product.status),
                product.description or "",
                _fmt_time(product.created_at),
            ]
        )

    wb.save(path)
    return path


def remove_file(path: str) -> None:
    try:
        os.remove(path)
    except OSError:
        pass


def build_filename(prefix: str) -> str:
    return f"{prefix}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
